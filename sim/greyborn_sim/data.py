"""Load every tuning value the simulator uses from spec/greyborn-tuning.xlsx.

The workbook is read with cached values (data_only=True), so it must have been
recalculated after its last edit (Excel does this on save; after an openpyxl
edit, run the xlsx skill's recalc.py).
"""
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

DEFAULT_WORKBOOK = Path(__file__).resolve().parents[2] / "spec" / "greyborn-tuning.xlsx"

LINEAGES = ["Titan", "Brawler", "Verdant", "Hollow", "Thornrunner", "Bonespire", "Stillheart"]
STATS = ["Health", "Power", "Speed %", "Control"]
PAIRS = [("Titan", "Brawler"), ("Verdant", "Hollow"), ("Thornrunner", "Bonespire")]

# Base Form basic attack (Spec 02, section 0): 1.0 x P per hit, 1.2 hits/s.
BASE_FORM_BASIC = (1.0, 1.2)


class WorkbookNotRecalculated(Exception):
    pass


@dataclass
class Tuning:
    stats: dict = field(default_factory=dict)        # lineage -> stat -> [None, L1..L20]
    exp_to_next: list = field(default_factory=list)  # index = level (1..19)
    exp_params: dict = field(default_factory=dict)
    abilities: dict = field(default_factory=dict)    # lineage -> list of dicts
    basic: dict = field(default_factory=dict)        # lineage -> (ratio, hits per second)
    ults: dict = field(default_factory=dict)         # lineage -> list of dicts
    synergies: list = field(default_factory=list)
    phase_synergy: dict = field(default_factory=dict)  # phase -> [S1, S2, S3] cooldown or None
    lullaby: float = 0.0
    creatures: list = field(default_factory=list)
    phase_mult: list = field(default_factory=list)   # index 0..3 for phases 1..4
    affinity_mult: float = 1.25
    econ: dict = field(default_factory=dict)
    tension: dict = field(default_factory=dict)
    events: list = field(default_factory=list)

    def stat(self, lineage, name, level):
        return self.stats[lineage][name][level]

    @property
    def exp_total(self):
        return sum(self.exp_to_next[1:20])


def _need(value, where):
    if value is None:
        raise WorkbookNotRecalculated(
            f"{where} has no cached value. Open and save the workbook in Excel or "
            "LibreOffice (or run the xlsx skill's recalc.py) before simulating.")
    return value


def _num(text):
    m = re.search(r"-?\d[\d,]*\.?\d*", str(text or ""))
    return float(m.group().replace(",", "")) if m else None


def load(path=DEFAULT_WORKBOOK, overrides=None):
    wb = load_workbook(path, data_only=True)
    t = Tuning()

    ws = wb["Stats by level"]
    for r in range(5, ws.max_row + 1):
        key = ws.cell(row=r, column=1).value
        if not key:
            continue
        lineage, stat = key.split("|")
        row = [None] + [_need(ws.cell(row=r, column=6 + lv).value, f"Stats by level!{key} L{lv}")
                        for lv in range(1, 21)]
        t.stats.setdefault(lineage, {})[stat] = row

    ws = wb["EXP & Respawn"]
    for r in range(4, 13):
        name = ws.cell(row=r, column=10).value
        if name:
            t.exp_params[name] = ws.cell(row=r, column=11).value
    t.exp_to_next = [None] + [_need(ws.cell(row=4 + lv, column=2).value, f"EXP & Respawn!B{4 + lv}")
                              for lv in range(1, 20)]

    ws = wb["Abilities"]
    for r in range(4, ws.max_row + 1):
        lineage, slot = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
        if lineage not in LINEAGES:
            continue
        base, ratio, cd = (ws.cell(row=r, column=c).value for c in (8, 9, 10))
        if slot == "Basic":
            t.basic[lineage] = (ratio, _num(ws.cell(row=r, column=5).value))
        else:
            t.abilities.setdefault(lineage, []).append(
                {"slot": slot, "name": ws.cell(row=r, column=3).value,
                 "base": base, "ratio": ratio, "cd": cd})

    ws = wb["Ultimates"]
    for r in range(4, ws.max_row + 1):
        lineage = ws.cell(row=r, column=1).value
        if lineage not in LINEAGES:
            continue
        t.ults.setdefault(lineage, []).append({
            "name": ws.cell(row=r, column=3).value,
            "base": ws.cell(row=r, column=8).value,
            "ratio": ws.cell(row=r, column=9).value,
            "cd": ws.cell(row=r, column=10).value,
        })

    ws = wb["Synergies"]
    for r in range(4, 7):
        pair = tuple(x.strip() for x in ws.cell(row=r, column=2).value.split("+"))
        t.synergies.append({
            "name": ws.cell(row=r, column=1).value, "pair": pair,
            "unlock": ws.cell(row=r, column=3).value,
            "cds": [ws.cell(row=r, column=c).value for c in (6, 7, 8)],
        })
    for i, r in enumerate(range(11, 15), start=1):
        vals = [ws.cell(row=r, column=c).value for c in (2, 3, 4)]
        t.phase_synergy[i] = [None if v == "Locked" else _need(v, f"Synergies!row {r}") for v in vals]
    t.lullaby = ws["E16"].value

    ws = wb["Creatures"]
    t.phase_mult = [ws[c].value for c in ("D2", "F2", "H2", "J2")]
    t.affinity_mult = ws["L2"].value
    for r in range(4, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        tier = ws.cell(row=r, column=4).value
        if not name or not tier:
            continue
        health = ws.cell(row=r, column=5).value
        t.creatures.append({
            "name": name, "group": ws.cell(row=r, column=3).value, "tier": str(tier).split()[0],
            "health": health if isinstance(health, (int, float)) else None,
            "cores": ws.cell(row=r, column=10).value or 0,
            "sap": ws.cell(row=r, column=11).value or 0,
        })

    ws = wb["Economy"]
    for r in range(4, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if name and ws.cell(row=r, column=2).value is not None:
            t.econ[name] = ws.cell(row=r, column=2).value

    ws = wb["Tension & Events"]
    for r in range(4, 11):
        t.tension[ws.cell(row=r, column=1).value] = ws.cell(row=r, column=2).value
    for r in range(18, ws.max_row + 1):
        deck = ws.cell(row=r, column=1).value
        if deck in ("Neutral", "Planet Pulse", "Murmur Surge"):
            t.events.append({"deck": deck, "name": ws.cell(row=r, column=2).value,
                             "weight": ws.cell(row=r, column=3).value,
                             "duration": ws.cell(row=r, column=4).value})

    for key, value in (overrides or {}).items():
        apply_override(t, key, value)
    return t


def apply_override(t, key, value):
    """Override one value: 'Economy:<parameter>', 'EXP:<input>', 'Tension:<factor>'."""
    sheet, _, name = key.partition(":")
    table = {"Economy": t.econ, "EXP": t.exp_params, "Tension": t.tension}.get(sheet)
    if table is None or name not in table:
        raise KeyError(f"Unknown override '{key}'. Use Economy:, EXP: or Tension: with an exact row name.")
    table[name] = value
    if sheet == "EXP":
        base, step = t.exp_params["EXP base"], t.exp_params["EXP per level"]
        t.exp_to_next = [None] + [base + step * lv for lv in range(1, 20)]
